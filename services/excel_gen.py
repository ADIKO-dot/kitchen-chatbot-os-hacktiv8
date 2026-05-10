"""
Excel Generation Service — Creates formatted spreadsheets for inventory and schedules.
"""

import os
from datetime import date
from io import BytesIO
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from loguru import logger

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "static", "generated")
os.makedirs(OUTPUT_DIR, exist_ok=True)


async def generate_inventory_excel(items: list[dict]) -> str:
    """
    Generate an inventory spreadsheet.
    items: [{"name": str, "qty": float, "unit": str, "cost_per_unit": float}]
    Returns: file path to the generated .xlsx
    """
    df = pd.DataFrame(items)
    df["total_value"] = df["qty"] * df["cost_per_unit"]

    filename = f"inventory_{date.today().isoformat()}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Inventory")
        ws = writer.sheets["Inventory"]
        # Style header row
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        # Auto-width columns
        for col in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 2

    logger.info(f"[ExcelGen] Created inventory: {filepath}")
    return filepath


async def generate_schedule_excel(entries: list[dict]) -> str:
    """
    Generate a staff schedule spreadsheet.
    entries: [{"staff": str, "date": str, "shift": str}]
    Returns: file path to the generated .xlsx
    """
    df = pd.DataFrame(entries)

    filename = f"schedule_{date.today().isoformat()}.xlsx"
    filepath = os.path.join(OUTPUT_DIR, filename)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Schedule")
        ws = writer.sheets["Schedule"]
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            max_len = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 2

    logger.info(f"[ExcelGen] Created schedule: {filepath}")
    return filepath
