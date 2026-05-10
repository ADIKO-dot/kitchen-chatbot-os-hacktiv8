"""
Document Loader — Parses uploaded files (PDF, DOCX, TXT, CSV, XLSX) into text chunks for RAG.
"""

import os
import csv
import io
from loguru import logger


def load_text(content: bytes) -> str:
    return content.decode("utf-8", errors="ignore")


def load_csv(content: bytes) -> str:
    text = content.decode("utf-8", errors="ignore")
    reader = csv.reader(io.StringIO(text))
    rows = [" | ".join(row) for row in reader]
    return "\n".join(rows)


def load_pdf(content: bytes) -> str:
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages)
    except ImportError:
        raise ImportError("PyPDF2 is required for PDF parsing. Install: pip install PyPDF2")


def load_docx(content: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except ImportError:
        raise ImportError("python-docx is required for DOCX parsing. Install: pip install python-docx")


def load_xlsx(content: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
        lines = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            lines.append(f"[Sheet: {sheet}]")
            for row in ws.iter_rows(values_only=True):
                lines.append(" | ".join(str(c) if c is not None else "" for c in row))
        return "\n".join(lines)
    except ImportError:
        raise ImportError("openpyxl is required for XLSX parsing. Install: pip install openpyxl")


LOADERS = {
    ".txt": load_text,
    ".csv": load_csv,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".xlsx": load_xlsx,
}

SUPPORTED_EXTENSIONS = list(LOADERS.keys())


def parse_file(filename: str, content: bytes) -> str:
    """Parse file content based on extension. Returns extracted text."""
    ext = os.path.splitext(filename)[1].lower()
    loader = LOADERS.get(ext)
    if not loader:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {SUPPORTED_EXTENSIONS}")
    text = loader(content)
    logger.info(f"[DocLoader] Parsed {filename} ({ext}): {len(text)} chars")
    return text


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 100) -> list[str]:
    """Split text into overlapping chunks using LangChain's RecursiveCharacterTextSplitter."""
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", "。", ".", "!", "?", ";", ",", " ", ""],
    )
    return splitter.split_text(text)
